#!/usr/bin/env python3
"""Costruzione del workbook Excel dell'inventario testi — Arcade Maze Fantasy.

Recovery Task R1: ricrea ArcadeMaze_Testi_Gioco.xlsx (originale Task 103,
perso nel reset sandbox) a partire da game_texts.json prodotto da
extract_game_texts.py.

Struttura (documentata nel §13 del contesto): 15 fogli — Menu, Selezione,
Intro, HUD-pausa, Level Up, Fine partita, Bonus reliquie, Card bonus, Armi,
Armature, Nemici-Boss, Crediti, Altro, Immagini, LEGGIMI. Colonne: ID
(T001+), Riferimento (chiave/const), Testo IT, Testo EN, Dove appare,
Posizione nel codice (file:riga), Note.

Nessuna formula (inventario puro) -> zero errori formula per costruzione.
Auto-verifica: riapre il file da disco e controlla fogli/righe/ID.

Uso:
  python3 build_texts_xlsx.py [--json game_texts.json]
                              [--out /home/z/my-project/download/ArcadeMaze_Testi_Gioco.xlsx]
"""
from __future__ import annotations

import argparse
import json
import os
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SHEET_ORDER = ["Menu", "Selezione", "Intro", "HUD-pausa", "Level Up",
               "Fine partita", "Bonus reliquie", "Card bonus", "Armi",
               "Armature", "Nemici-Boss", "Crediti", "Altro", "Immagini"]

HEADERS = ["ID", "Riferimento", "Testo IT", "Testo EN",
           "Dove appare", "Posizione nel codice", "Note"]
COL_WIDTHS = [9, 34, 48, 48, 30, 44, 42]

# ---- design token (tema oro/scuro medievale, coerente col progetto) -------
C_TITLE_BG = "2B1D0E"   # bruno scurissimo
C_TITLE_FG = "E3C05A"   # oro chiaro
C_HEAD_BG = "3B2A1A"    # bruno scuro
C_HEAD_FG = "FFFFFF"
C_ROW_ALT = "F7F1E3"    # pergamena chiara
C_BORDER = "B9A88C"
C_ID_FG = "7A5C14"      # oro scuro

F_TITLE = Font(name="Calibri", size=14, bold=True, color=C_TITLE_FG)
F_HEAD = Font(name="Calibri", size=11, bold=True, color=C_HEAD_FG)
F_BODY = Font(name="Calibri", size=11, color="1A1A1A")
F_ID = Font(name="Calibri", size=11, bold=True, color=C_ID_FG)

FILL_TITLE = PatternFill("solid", fgColor=C_TITLE_BG)
FILL_HEAD = PatternFill("solid", fgColor=C_HEAD_BG)
FILL_ALT = PatternFill("solid", fgColor=C_ROW_ALT)

THIN = Side(style="thin", color=C_BORDER)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

AL_LEFT = Alignment(horizontal="left", vertical="top", wrap_text=True)
AL_CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)


def safe_text(v) -> str:
    """None -> '' · testo che inizia con '=' -> prefisso apostrofo (skill rule)."""
    if v is None:
        return ""
    s = str(v)
    return "'" + s if s.startswith("=") else s


def build_data_sheet(ws, sheet_title: str, rows: list[dict]):
    ncols = len(HEADERS)
    # riga 1: titolo (merged)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(row=1, column=1,
                value=f"ARCADE MAZE FANTASY — {sheet_title} · Inventario testi")
    c.font = F_TITLE
    c.fill = FILL_TITLE
    c.alignment = Alignment(horizontal="left", vertical="center")
    for col in range(1, ncols + 1):
        ws.cell(row=1, column=col).fill = FILL_TITLE
    ws.row_dimensions[1].height = 26
    # riga 2: header
    for col, h in enumerate(HEADERS, start=1):
        c = ws.cell(row=2, column=col, value=h)
        c.font = F_HEAD
        c.fill = FILL_HEAD
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[2].height = 22
    # dati
    for i, e in enumerate(rows):
        r = i + 3
        values = [safe_text(e.get("id")), safe_text(e.get("rif")),
                  safe_text(e.get("it")), safe_text(e.get("en")),
                  safe_text(e.get("dove")), safe_text(e.get("pos")),
                  safe_text(e.get("note"))]
        for col, v in enumerate(values, start=1):
            c = ws.cell(row=r, column=col, value=v)
            c.border = BORDER
            c.alignment = AL_CENTER if col == 1 else AL_LEFT
            c.font = F_ID if col == 1 else F_BODY
            if i % 2 == 1:
                c.fill = FILL_ALT
    # larghezze colonne
    for col, w in enumerate(COL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col)].width = w
    # freeze + filtri
    ws.freeze_panes = "A3"
    last = max(2, len(rows) + 2)
    ws.auto_filter.ref = f"A2:{get_column_letter(ncols)}{last}"


def build_leggimi(ws, meta: dict, per_sheet: dict, total: int):
    ws.sheet_properties.tabColor = "C9A227"
    rows = [
        ("Cos'e' questo file",
         "Inventario completo dei testi visibili nel gioco (IT ed EN) per la revisione "
         "linguistica. Ogni riga = una stringa mostrata a schermo, con la posizione "
         "esatta nel codice (file:riga) e dove appare in gioco."),
        ("RIGENERAZIONE post-reset (importante)",
         f"Questo file e' stato RIGENERATO il {meta.get('generated', '?')} (Task R1) "
         "dopo un reset sandbox che ha cancellato l'originale del Task 103 e i suoi "
         "script generatori. Totale: 438 voci (l'originale ne aveva 439). Gli ID "
         "T001-T438 sono stati riassegnati e POSSONO non coincidere con gli ID di "
         "una copia originale scaricata in precedenza: se hai annotazioni su una "
         "copia vecchia, cita SEMPRE anche la colonna 'Riferimento' (chiave/costante) "
         "oltre all'ID."),
        ("Come richiedere una correzione",
         "Cita l'ID (o il Riferimento) + il nuovo Testo IT + il nuovo Testo EN. "
         "Destinazioni delle correzioni: I18nManager.gd TEXTS (291 chiavi), "
         "GameConstants.gd (SKILL_DEFS / RELIC_* / ELEMENTAL_WEAPONS), "
         "Boss.gd get_boss_name() (17 nomi), stringhe hardcoded nei file ui, "
         "PNG testuali (richiedono rigenerazione immagine, non modifica testo)."),
        ("Anomalie note (revisione Task 103)",
         "• GAME OVER hardcoded in LoseScreen.gd:233 (resta in inglese anche in IT)\n"
         "• Crediti: ruoli 'Luca A. Greco' e publisher 'Marled Software' hardcoded "
         "(CreditsScreen.gd:229/221/323)\n"
         "• Footer menu 'By' + 'Marled Software' hardcoded (MainMenu.gd:610/629)\n"
         "• 'P1'/'P2' hardcoded nell'HUD 2P (HUD.gd:405/430)\n"
         "• Chiavi orfane: hud.energy, hud.en (barra rimossa Task 96), "
         "menu.lang_english, menu.lang_italian, enemy.generic, cr.title\n"
         "• Doppia definizione: le 17 chiavi boss.* di TEXTS NON sono usate; la "
         "fonte reale dei nomi boss e' Boss.get_boss_name()\n"
         "• Short reliquia (nome corto) MAI mostrati a schermo\n"
         "• Nome/descrizione armi elementali MAI mostrati a schermo\n"
         "• Modale Level Up: nome IT grande + EN piccolo insieme; descrizione "
         "sempre italiana\n"
         "• PNG orfani non referenziati dal codice: header_credits.png, "
         "header_gameover.png, header_relic.png"),
        ("Fogli e provenienza",
         "Menu (menu.*, mode.* + footer) · Selezione (char.*, select.*) · Intro "
         "(intro.*) · HUD-pausa (hud.*, pause.* + P1/P2) · Level Up (levelup.* + "
         "SKILL_DEFS) · Fine partita (win/lose/cont/hof/ne/screen + GAME OVER) · "
         "Bonus reliquie (bonus.* + RELIC_MELEE/RELIC_RANGED) · Card bonus "
         "(bonuscard.*) · Armi (get_weapon_name + ELEMENTAL_WEAPONS) · Armature "
         "(RELIC_ARMORS) · Nemici-Boss (enemy.*, boss.*, weakness + get_boss_name) "
         "· Crediti (cr.* + hardcoded) · Altro (joy/demo/loading/misc) · Immagini "
         "(PNG testuali) · LEGGIMI (questo foglio)."),
        ("Statistiche",
         " · ".join(f"{s}: {per_sheet.get(s, 0)}" for s in SHEET_ORDER)
         + f" · TOTALE: {total}"),
        ("Rigenerare questo file in futuro",
         "Gli script generatori sono ora COMMITTATI nel repo (lezione 50: cio' che "
         "deve sopravvivere va nel repo): ArcadeMaze/scripts/extract_game_texts.py "
         "produce game_texts.json; ArcadeMaze/scripts/build_texts_xlsx.py produce "
         "questo workbook. Nessun credito AI richiesto: estrazione statica."),
        ("Convenzioni colonne",
         "ID = identificativo univoco (T001+). Riferimento = chiave TEXTS, costante "
         "o funzione nel codice. Testo IT/EN = testo attualmente a schermo (i "
         "segnaposto %d/%s sono formati applicati a runtime). 'Dove appare' e' "
         "euristico per area (rigenerato). Posizione = file:riga di definizione + "
         "primi utilizzi trovati."),
    ]
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 110
    # riga titolo
    ws.merge_cells("A1:B1")
    c = ws.cell(row=1, column=1, value="LEGGIMI — Arcade Maze Fantasy · Inventario testi")
    c.font = F_TITLE
    c.fill = FILL_TITLE
    ws.cell(row=1, column=2).fill = FILL_TITLE
    ws.row_dimensions[1].height = 26
    for col, h in enumerate(["Argomento", "Dettaglio"], start=1):
        c = ws.cell(row=2, column=col, value=h)
        c.font = F_HEAD
        c.fill = FILL_HEAD
        c.border = BORDER
        c.alignment = Alignment(horizontal="left", vertical="center")
    for i, (k, v) in enumerate(rows):
        r = i + 3
        a = ws.cell(row=r, column=1, value=k)
        b = ws.cell(row=r, column=2, value=v)
        a.font = Font(name="Calibri", size=11, bold=True, color=C_ID_FG)
        b.font = F_BODY
        for c2 in (a, b):
            c2.alignment = AL_LEFT
            c2.border = BORDER
            if i % 2 == 1:
                c2.fill = FILL_ALT
    ws.freeze_panes = "A3"


def main() -> int:
    here = os.path.dirname(os.path.abspath(__file__))
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--json", default=os.path.join(here, "game_texts.json"))
    default_out = ("/home/z/my-project/download/ArcadeMaze_Testi_Gioco.xlsx"
                   if os.path.isdir("/home/z/my-project/download")
                   else "ArcadeMaze_Testi_Gioco.xlsx")
    ap.add_argument("--out", default=default_out)
    args = ap.parse_args()

    with open(args.json, encoding="utf-8") as f:
        data = json.load(f)
    entries = data["entries"]

    # assegna gli ID T001+ nell'ordine (foglio, ordine di emissione)
    by_sheet: dict[str, list[dict]] = {s: [] for s in SHEET_ORDER}
    unknown = []
    for e in entries:
        s = e.get("sheet")
        if s in by_sheet:
            by_sheet[s].append(e)
        else:
            unknown.append(e)
    if unknown:
        print(f"ERRORE: {len(unknown)} voci con foglio sconosciuto: "
              f"{sorted({e['sheet'] for e in unknown})}", file=sys.stderr)
        return 2
    n = 0
    for s in SHEET_ORDER:
        for e in by_sheet[s]:
            n += 1
            e["id"] = f"T{n:03d}"

    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.creator = "Z.ai"
    wb.properties.title = "Arcade Maze Fantasy — Inventario testi (revisione IT/EN)"

    for s in SHEET_ORDER:
        ws = wb.create_sheet(title=s)
        build_data_sheet(ws, s, by_sheet[s])
    build_leggimi(wb.create_sheet(title="LEGGIMI"), data,
                  {s: len(by_sheet[s]) for s in SHEET_ORDER}, n)

    os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
    wb.save(args.out)
    print(f"salvato: {args.out}")

    # ------------------------------------------------ auto-verifica (disk) --
    wb2 = load_workbook(args.out, read_only=True)
    sheets = wb2.sheetnames
    assert sheets == SHEET_ORDER + ["LEGGIMI"], f"fogli errati: {sheets}"
    problems = []
    expected_id = 0
    for s in SHEET_ORDER:
        ws2 = wb2[s]
        rows = list(ws2.iter_rows(min_row=3, max_col=7, values_only=True))
        data_rows = [r for r in rows if r[0] is not None]
        if len(data_rows) != len(by_sheet[s]):
            problems.append(f"{s}: {len(data_rows)} righe su disco vs {len(by_sheet[s])}")
        for r in data_rows:
            expected_id += 1
            if r[0] != f"T{expected_id:03d}":
                problems.append(f"{s}: ID fuori sequenza: {r[0]} (atteso T{expected_id:03d})")
                expected_id = int(str(r[0])[1:])
            if not r[1]:
                problems.append(f"{s} {r[0]}: Riferimento vuoto")
            if not r[2] and not r[3]:
                problems.append(f"{s} {r[0]}: Testo IT e EN entrambi vuoti")
    wb2.close()
    if problems:
        print("VERIFICA FALLITA:")
        for p in problems[:20]:
            print("  !", p)
        return 1
    print(f"verifica OK: 15 fogli · {expected_id} voci · sequenza ID T001-T{expected_id:03d} "
          f"· nessun campo obbligatorio vuoto")
    return 0


if __name__ == "__main__":
    sys.exit(main())
